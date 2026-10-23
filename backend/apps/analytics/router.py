import io
from typing import Annotated

from dateutil import parser as date_parser
from fastapi import APIRouter, Depends, Response
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.analytics.repository import build_summary
from core.deps import MongoUser, require_support_staff

router = APIRouter(prefix="/api/analytics", tags=["analytics"])


def _parse_date(value):
    if not value:
        return None
    try:
        return date_parser.parse(value)
    except (ValueError, OverflowError):
        return None


@router.get("/summary", response_model=None)
def analytics_summary(
    _staff: Annotated[MongoUser, Depends(require_support_staff)],
    date_from: str | None = None,
    date_to: str | None = None,
):
    return build_summary(_parse_date(date_from), _parse_date(date_to))


@router.get("/export/excel", response_model=None)
def analytics_export_excel(
    _staff: Annotated[MongoUser, Depends(require_support_staff)],
    date_from: str | None = None,
    date_to: str | None = None,
):
    summary = build_summary(_parse_date(date_from), _parse_date(date_to))

    wb = Workbook()

    overview = wb.active
    overview.title = "Overview"
    overview.append(["Metric", "Value"])
    overview.append(["Total Tickets", summary["total_tickets"]])
    overview.append(["Avg. Resolution Time (hrs)", summary["avg_resolution_time_hours"]])
    overview.append(["Avg. First Response Time (min)", summary["avg_first_response_time_minutes"]])
    overview.append(["Reopened Tickets", summary["reopened_tickets"]])
    overview.append(["Pending Tickets", summary["pending_tickets"]])
    overview.append(["CSAT %", summary["csat"]["csat_percentage"]])
    overview.append(["SLA Compliance %", summary["sla_compliance"]["sla_percentage"]])

    by_status = wb.create_sheet("By Status")
    by_status.append(["Status", "Count"])
    for status, count in summary["tickets_by_status"].items():
        by_status.append([status, count])

    by_platform = wb.create_sheet("By Platform")
    by_platform.append(["Platform", "Count"])
    for platform, count in summary["tickets_by_platform"].items():
        by_platform.append([platform, count])

    by_user_type = wb.create_sheet("By User Type")
    by_user_type.append(["User Type", "Count"])
    for user_type, count in summary["tickets_by_user_type"].items():
        by_user_type.append([user_type, count])

    per_engineer = wb.create_sheet("Resolved Per Engineer")
    per_engineer.append(["Engineer", "Resolved Count"])
    for row in summary["resolved_per_engineer"]:
        per_engineer.append([row["engineer"], row["resolved_count"]])

    trend = wb.create_sheet("Daily Trend")
    trend.append(["Date", "Ticket Count"])
    for row in summary["daily_trend"]:
        trend.append([row["period"], row["count"]])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return Response(
        content=buffer.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="tech_support_report.xlsx"'},
    )


def _styled_table(rows):
    table = Table(rows, colWidths=[8 * cm, 6 * cm])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
            ]
        )
    )
    return table


@router.get("/export/pdf", response_model=None)
def analytics_export_pdf(
    _staff: Annotated[MongoUser, Depends(require_support_staff)],
    date_from: str | None = None,
    date_to: str | None = None,
):
    summary = build_summary(_parse_date(date_from), _parse_date(date_to))

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = [Paragraph("GoMechanic Tech Support — Analytics Report", styles["Title"]), Spacer(1, 12)]

    overview_rows = [
        ["Metric", "Value"],
        ["Total Tickets", summary["total_tickets"]],
        ["Avg. Resolution Time (hrs)", summary["avg_resolution_time_hours"]],
        ["Avg. First Response Time (min)", summary["avg_first_response_time_minutes"]],
        ["Reopened Tickets", summary["reopened_tickets"]],
        ["Pending Tickets", summary["pending_tickets"]],
        ["CSAT %", summary["csat"]["csat_percentage"]],
        ["SLA Compliance %", summary["sla_compliance"]["sla_percentage"]],
    ]
    elements.append(_styled_table(overview_rows))
    elements.append(Spacer(1, 16))

    elements.append(Paragraph("Tickets by Status", styles["Heading2"]))
    elements.append(_styled_table([["Status", "Count"], *summary["tickets_by_status"].items()]))
    elements.append(Spacer(1, 16))

    elements.append(Paragraph("Tickets by Platform", styles["Heading2"]))
    elements.append(_styled_table([["Platform", "Count"], *summary["tickets_by_platform"].items()]))
    elements.append(Spacer(1, 16))

    elements.append(Paragraph("Resolved Per Engineer", styles["Heading2"]))
    elements.append(
        _styled_table(
            [["Engineer", "Resolved Count"]]
            + [[row["engineer"], row["resolved_count"]] for row in summary["resolved_per_engineer"]]
        )
    )

    doc.build(elements)
    buffer.seek(0)

    return Response(
        content=buffer.read(),
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="tech_support_report.pdf"'},
    )
